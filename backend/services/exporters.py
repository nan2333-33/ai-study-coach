"""周报导出：把结构化文档渲染成 Markdown / Word / PDF。

调用方是 routes/report.py，输入统一为 `_report_doc()` 产出的结构化文档：

    {
      "title": "学习周报（2026-10-12 ~ 2026-10-18）",
      "blocks": [
        {"type": "h2",    "text": "做题量"},
        {"type": "p",     "text": "本周共答 42 题……"},
        {"type": "ol",    "items": ["...", "..."]},
        {"type": "quote", "text": "继续保持节奏"}
      ]
    }

三种格式共用同一份结构，避免 Markdown 改了 Word 没跟着改。
"""

import re
import io

# ---------------------------------------------------------------- 行内加粗

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def parse_inline(text):
    """把 `**强调**` 解析成 [(文本, 是否加粗), ...]，供 Word / PDF 渲染。"""
    segments = []
    pos = 0
    for m in _BOLD_RE.finditer(text or ""):
        if m.start() > pos:
            segments.append((text[pos : m.start()], False))
        segments.append((m.group(1), True))
        pos = m.end()
    if pos < len(text or ""):
        segments.append((text[pos:], False))
    return segments or [(text or "", False)]


# ---------------------------------------------------------------- Markdown


def to_markdown(doc):
    lines = [f"# {doc['title']}", ""]
    for b in doc["blocks"]:
        t = b["type"]
        if t == "h2":
            lines += [f"## {b['text']}", ""]
        elif t == "p":
            lines += [b["text"], ""]
        elif t == "ol":
            lines += [f"{i}. {x}" for i, x in enumerate(b["items"], 1)]
            lines.append("")
        elif t == "quote":
            lines += [f"> {b['text']}", ""]
    return "\n".join(lines)


# ---------------------------------------------------------------- Word

CJK_FONT = "Microsoft YaHei"


def to_docx(doc):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = Document()

    # 正文字体改成微软雅黑：只设 font.name 对中文无效，必须同时设 w:eastAsia
    normal = document.styles["Normal"]
    normal.font.name = CJK_FONT
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    rfonts.set(qn("w:ascii"), CJK_FONT)
    rfonts.set(qn("w:hAnsi"), CJK_FONT)

    def add_run(paragraph, text, bold=False, size=None):
        run = paragraph.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        run.font.name = CJK_FONT
        run._element.get_or_add_rPr()
        rf = run._element.rPr.find(qn("w:rFonts"))
        if rf is None:
            from docx.oxml import OxmlElement

            rf = OxmlElement("w:rFonts")
            run._element.rPr.append(rf)
        rf.set(qn("w:eastAsia"), CJK_FONT)
        return run

    document.add_heading(doc["title"], level=0)
    for b in doc["blocks"]:
        t = b["type"]
        if t == "h2":
            document.add_heading(b["text"], level=2)
        elif t == "p":
            para = document.add_paragraph()
            for text, bold in parse_inline(b["text"]):
                add_run(para, text, bold=bold)
        elif t == "ol":
            for item in b["items"]:
                para = document.add_paragraph(style="List Number")
                for text, bold in parse_inline(item):
                    add_run(para, text, bold=bold)
        elif t == "quote":
            para = document.add_paragraph()
            para.paragraph_format.left_indent = Pt(24)
            for text, bold in parse_inline(b["text"]):
                add_run(para, text, bold=bold, size=11)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

_CJK_PDF_FONTS = None


def _cjk_pdf_fonts():
    """注册并返回 (常规字体名, 粗体字体名)，按 Windows → macOS → reportlab 内置 CID 依次回退。"""
    global _CJK_PDF_FONTS
    if _CJK_PDF_FONTS:
        return _CJK_PDF_FONTS

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    candidates = [
        ("CJK-Regular", "CJK-Bold", "C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/msyhbd.ttc"),
        ("CJK-Regular", "CJK-Bold", "/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/PingFang.ttc"),
        ("CJK-Regular", "CJK-Bold", "/usr/share/fonts/truetype/arphic/uming.ttc", None),
    ]
    for reg_name, bold_name, reg_path, bold_path in candidates:
        import os

        if not os.path.exists(reg_path):
            continue
        try:
            pdfmetrics.registerFont(TTFont(reg_name, reg_path, subfontIndex=0))
            if bold_path and os.path.exists(bold_path):
                pdfmetrics.registerFont(TTFont(bold_name, bold_path, subfontIndex=0))
            else:
                bold_name = reg_name
            from reportlab.lib.fonts import addMapping

            addMapping(reg_name, 0, 0, reg_name)
            addMapping(reg_name, 1, 0, bold_name)
            _CJK_PDF_FONTS = (reg_name, bold_name)
            return _CJK_PDF_FONTS
        except Exception:
            continue

    # 兜底：reportlab 自带的 Adobe 中文 CID 字体，不需要系统字体文件
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    _CJK_PDF_FONTS = ("STSong-Light", "STSong-Light")
    return _CJK_PDF_FONTS


def _xml_escape(text):
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def to_pdf(doc):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    regular, bold = _cjk_pdf_fonts()

    title_style = ParagraphStyle(
        "TitleCJK", fontName=bold, fontSize=18, leading=26, spaceAfter=10,
        textColor=colors.HexColor("#1f2937"),
    )
    h2_style = ParagraphStyle(
        "H2CJK", fontName=bold, fontSize=13, leading=20, spaceBefore=12, spaceAfter=4,
        textColor=colors.HexColor("#4f46e5"),
    )
    body_style = ParagraphStyle(
        "BodyCJK", fontName=regular, fontSize=11, leading=18, spaceAfter=4,
        textColor=colors.HexColor("#374151"),
    )
    quote_style = ParagraphStyle(
        "QuoteCJK", fontName=regular, fontSize=11, leading=18, spaceAfter=4,
        leftIndent=12, textColor=colors.HexColor("#6b7280"),
    )

    def rich(text):
        out = []
        for seg, is_bold in parse_inline(text):
            seg = _xml_escape(seg)
            out.append(f"<b>{seg}</b>" if is_bold else seg)
        return "".join(out)

    buf = io.BytesIO()
    document = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
        title=doc["title"],
    )
    story = [Paragraph(_xml_escape(doc["title"]), title_style)]
    for b in doc["blocks"]:
        t = b["type"]
        if t == "h2":
            story.append(Paragraph(_xml_escape(b["text"]), h2_style))
        elif t == "p":
            story.append(Paragraph(rich(b["text"]), body_style))
        elif t == "ol":
            for i, item in enumerate(b["items"], 1):
                story.append(Paragraph(f"{i}. {rich(item)}", body_style))
        elif t == "quote":
            story.append(Paragraph(rich(b["text"]), quote_style))
        story.append(Spacer(1, 2))
    document.build(story)
    return buf.getvalue()
