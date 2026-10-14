"""文本提取：从 PDF / PPTX / DOCX / MD / TXT 中读取文字，并保留页码（段落级来源）。

返回结构：List[PageText]，page_number 从 1 开始，text 为该页文字。
- PDF     : 逐页提取（pypdf）
- PPTX    : 逐页提取文本 + 备注（python-pptx）
- DOCX    : 逐段提取（python-docx），整篇作为一页流
- MD / TXT: 直接读取全文，视为单页

对应 PRD M-02（文本提取）与 M-05（来源元数据：文件名 / 页码）。
"""
from dataclasses import dataclass
from pathlib import Path


@dataclass
class PageText:
    page_number: int
    text: str


def extract_text(file_path: str, file_type: str) -> list[PageText]:
    ft = (file_type or "").lower()
    if ft == "pdf":
        return _extract_pdf(file_path)
    if ft in ("ppt", "pptx"):
        return _extract_pptx(file_path)
    if ft in ("doc", "docx"):
        return _extract_docx(file_path)
    if ft in ("md", "markdown", "txt"):
        return _extract_plain(file_path)
    raise ValueError(f"不支持的文件类型: {file_type}")


def _extract_pdf(path: str) -> list[PageText]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(PageText(i, text))
    return pages


def _extract_pptx(path: str) -> list[PageText]:
    from pptx import Presentation

    prs = Presentation(path)
    pages = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes:
                parts.append(notes)
        pages.append(PageText(i, "\n".join(p for p in parts if p)))
    return pages


def _extract_docx(path: str) -> list[PageText]:
    from docx import Document

    doc = Document(path)
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n".join(paras)
    return [PageText(1, text)] if text else []


def _extract_plain(path: str) -> list[PageText]:
    text = Path(path).read_text(encoding="utf-8", errors="ignore")
    return [PageText(1, text)] if text.strip() else []
